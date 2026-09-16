"""Tests for Foundation Enrolment Pipeline endpoint (spec/03 §2, §5, FND-014)."""
import csv
from datetime import date, datetime, timedelta, timezone as dt_timezone
from decimal import Decimal
import io
import openpyxl
from rest_framework import status
from rest_framework.test import APITestCase
from django.utils import timezone

from apps.academic.models import AcademicYear, ClassGroup, ClassEnrollment
from apps.core.models import DomainEvent
from apps.foundation.services import get_foundation_enrolment_pipeline
from apps.identity.models import Foundation, School, Person, Student, User
from apps.identity.rbac import (
    assign_role, ROLE_FOUNDATION_ADMIN, ROLE_PARENT, SCOPE_FOUNDATION
)
from educore.middleware.tenancy import clear_current_foundation_id, tenant_context


class FoundationEnrolmentPipelineTests(APITestCase):
    def setUp(self):
        clear_current_foundation_id()

        # 1. Foundation Alpha
        self.foundation_alpha = Foundation.objects.create(
            legal_name="Yayasan Alpha Nusantara",
            brand_name="Alpha Nusantara",
            npwp="01.111.222.3-001.000",
            reporting_currency="IDR",
        )
        self.school_a1 = School.all_tenants.create(
            foundation_id=self.foundation_alpha.id,
            name="SD Alpha Satu",
            npsn="10100001",
            level=School.LEVEL_SD,
            base_currency="IDR",
            is_active=True,
        )
        self.school_a2 = School.all_tenants.create(
            foundation_id=self.foundation_alpha.id,
            name="SMP Alpha Dua",
            npsn="10100002",
            level=School.LEVEL_SMP,
            base_currency="IDR",
            is_active=True,
        )

        # 2. Foundation Beta (for cross-tenant isolation testing)
        self.foundation_beta = Foundation.objects.create(
            legal_name="Yayasan Beta Cendekia",
            brand_name="Beta Cendekia",
            reporting_currency="IDR",
        )
        self.school_b1 = School.all_tenants.create(
            foundation_id=self.foundation_beta.id,
            name="SMA Beta Cendekia",
            npsn="20200001",
            level=School.LEVEL_SMA,
            base_currency="IDR",
            is_active=True,
        )

        # 3. Users & Roles
        self.admin_alpha = User.all_tenants.create_user(
            phone_e164="+6281111111111",
            foundation_id=self.foundation_alpha.id,
            full_name="Ketua Yayasan Alpha",
        )
        assign_role(
            user=self.admin_alpha,
            role=ROLE_FOUNDATION_ADMIN,
            scope_type=SCOPE_FOUNDATION,
            scope_id=self.foundation_alpha.id,
            foundation_id=self.foundation_alpha.id,
        )

        self.user_no_perm = User.all_tenants.create_user(
            phone_e164="+6281222222222",
            foundation_id=self.foundation_alpha.id,
            full_name="Wali Murid Alpha",
        )
        assign_role(
            user=self.user_no_perm,
            role=ROLE_PARENT,
            scope_type=SCOPE_FOUNDATION,
            scope_id=self.foundation_alpha.id,
            foundation_id=self.foundation_alpha.id,
        )

        # 4. Academic Year & Classes for Alpha
        with tenant_context(self.foundation_alpha.id):
            self.ay_a = AcademicYear.objects.create(
                foundation_id=self.foundation_alpha.id,
                school=self.school_a1,
                name="2025/2026",
                start_date=date(2025, 7, 1),
                end_date=date(2026, 6, 30),
                is_active=True,
            )
            self.class_g1 = ClassGroup.objects.create(
                foundation_id=self.foundation_alpha.id,
                school=self.school_a1,
                academic_year=self.ay_a,
                grade_level=1,
                name="Kelas 1-A",
            )
            self.class_g2 = ClassGroup.objects.create(
                foundation_id=self.foundation_alpha.id,
                school=self.school_a1,
                academic_year=self.ay_a,
                grade_level=2,
                name="Kelas 2-A",
            )

            # 5. Create Persons & Students for Alpha
            # Student 1: Prospect Grade 1
            self.p1 = Person.objects.create(
                foundation_id=self.foundation_alpha.id,
                full_name="Budi Calon",
            )
            self.s1 = Student.objects.create(
                foundation_id=self.foundation_alpha.id,
                school=self.school_a1,
                person=self.p1,
                nis="NIS001",
                status=Student.STATUS_PROSPECT,
                target_grade_level=1,
            )

            # Student 2: Accepted Grade 1 in period (status ACTIVE, enrolled in period)
            self.p2 = Person.objects.create(
                foundation_id=self.foundation_alpha.id,
                full_name="Siti Diterima",
            )
            self.s2 = Student.objects.create(
                foundation_id=self.foundation_alpha.id,
                school=self.school_a1,
                person=self.p2,
                nis="NIS002",
                status=Student.STATUS_ACTIVE,
            )
            self.enr2 = ClassEnrollment.objects.create(
                foundation_id=self.foundation_alpha.id,
                student=self.s2,
                class_group=self.class_g1,
                enrolled_at=date(2025, 8, 10),
                is_active=True,
            )

            # Student 3: Active Grade 2 (enrolled prior, still active)
            self.p3 = Person.objects.create(
                foundation_id=self.foundation_alpha.id,
                full_name="Joko Aktif",
            )
            self.s3 = Student.objects.create(
                foundation_id=self.foundation_alpha.id,
                school=self.school_a1,
                person=self.p3,
                nis="NIS003",
                status=Student.STATUS_ACTIVE,
            )
            self.enr3 = ClassEnrollment.objects.create(
                foundation_id=self.foundation_alpha.id,
                student=self.s3,
                class_group=self.class_g2,
                enrolled_at=date(2025, 7, 5),
                is_active=True,
            )

            # Student 4: Churned (transferred out / inactive with status change event)
            self.p4 = Person.objects.create(
                foundation_id=self.foundation_alpha.id,
                full_name="Dewi Keluar",
            )
            self.s4 = Student.objects.create(
                foundation_id=self.foundation_alpha.id,
                school=self.school_a1,
                person=self.p4,
                nis="NIS004",
                status=Student.STATUS_TRANSFERRED_OUT,
                target_grade_level=2,
            )
            self.ev4 = DomainEvent.objects.create(
                foundation_id=self.foundation_alpha.id,
                name='identity.student.status_changed',
                payload={
                    'student_id': self.s4.id,
                    'school_id': self.school_a1.id,
                    'old_status': Student.STATUS_ACTIVE,
                    'new_status': Student.STATUS_TRANSFERRED_OUT,
                }
            )
            DomainEvent.objects.filter(id=self.ev4.id).update(occurred_at=datetime(2025, 8, 20, 10, 0, tzinfo=dt_timezone.utc))

            # Explicitly backdate created_at/updated_at for deterministic period testing
            Student.all_tenants.filter(id=self.s1.id).update(created_at=datetime(2025, 8, 1, 10, 0, tzinfo=dt_timezone.utc))
            Student.all_tenants.filter(id=self.s2.id).update(created_at=datetime(2025, 8, 10, 10, 0, tzinfo=dt_timezone.utc))
            Student.all_tenants.filter(id=self.s3.id).update(created_at=datetime(2025, 7, 5, 10, 0, tzinfo=dt_timezone.utc))
            Student.all_tenants.filter(id=self.s4.id).update(
                created_at=datetime(2025, 7, 1, 10, 0, tzinfo=dt_timezone.utc),
                updated_at=datetime(2025, 8, 20, 10, 0, tzinfo=dt_timezone.utc),
            )

        # 6. Foundation Beta Student (Cross-tenant leak check)
        with tenant_context(self.foundation_beta.id):
            p_beta = Person.objects.create(
                foundation_id=self.foundation_beta.id,
                full_name="Beta Student",
            )
            Student.objects.create(
                foundation_id=self.foundation_beta.id,
                school=self.school_b1,
                person=p_beta,
                nis="NISBETA1",
                status=Student.STATUS_PROSPECT,
                target_grade_level=1,
            )

        clear_current_foundation_id()

    def test_service_pipeline_metrics(self):
        result = get_foundation_enrolment_pipeline(
            foundation_id=self.foundation_alpha.id,
            from_date=date(2025, 8, 1),
            to_date=date(2025, 8, 31),
            group_by="campus"
        )
        summary = result["summary"]
        self.assertEqual(summary["total_schools"], 2)
        self.assertEqual(summary["total_prospects"], 1)
        self.assertEqual(summary["total_accepted"], 1)
        self.assertEqual(summary["total_active"], 2)  # s2 and s3
        self.assertEqual(summary["total_churned"], 1)  # s4
        # Conversion: 1 / (1 + 1) * 100 = 50.0%
        self.assertEqual(summary["conversion_rate_pct"], 50.0)
        # Retention: 2 / (2 + 1) * 100 = 66.67%
        self.assertEqual(summary["retention_rate_pct"], 66.67)

        # School A1 breakdown
        data = result["data"]
        a1_data = next((d for d in data if d["school_id"] == self.school_a1.id), None)
        self.assertIsNotNone(a1_data)
        self.assertEqual(a1_data["prospects"], 1)
        self.assertEqual(a1_data["accepted"], 1)
        self.assertEqual(a1_data["active"], 2)
        self.assertEqual(a1_data["churned"], 1)
        self.assertEqual(a1_data["conversion_rate_pct"], 50.0)
        self.assertEqual(a1_data["retention_rate_pct"], 66.67)

        # Grades inside School A1
        grades = {g["grade_level"]: g for g in a1_data["grades"]}
        self.assertIn(1, grades)
        self.assertEqual(grades[1]["prospects"], 1)
        self.assertEqual(grades[1]["accepted"], 1)
        self.assertEqual(grades[1]["active"], 1)
        self.assertIn(2, grades)
        self.assertEqual(grades[2]["active"], 1)
        self.assertEqual(grades[2]["churned"], 1)

    def test_service_pipeline_grouped_by_grade(self):
        result = get_foundation_enrolment_pipeline(
            foundation_id=self.foundation_alpha.id,
            from_date=date(2025, 8, 1),
            to_date=date(2025, 8, 31),
            group_by="grade"
        )
        self.assertEqual(result["group_by"], "grade")
        data = result["data"]
        g1 = next((d for d in data if d["grade_level"] == 1), None)
        self.assertIsNotNone(g1)
        self.assertEqual(g1["prospects"], 1)
        self.assertEqual(g1["accepted"], 1)
        self.assertEqual(g1["active"], 1)
        self.assertEqual(g1["churned"], 0)
        self.assertEqual(g1["conversion_rate_pct"], 50.0)
        self.assertEqual(g1["retention_rate_pct"], 100.0)

    def test_cross_tenant_isolation(self):
        result = get_foundation_enrolment_pipeline(
            foundation_id=self.foundation_alpha.id,
            from_date=date(2025, 1, 1),
            to_date=date(2025, 12, 31),
            group_by="campus"
        )
        for school_item in result["data"]:
            self.assertNotEqual(school_item["school_id"], self.school_b1.id)

    def test_endpoint_unauthenticated(self):
        response = self.client.get('/api/v1/foundation/enrolment')
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_endpoint_permission_denied(self):
        self.client.force_authenticate(user=self.user_no_perm)
        response = self.client.get('/api/v1/foundation/enrolment')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_endpoint_success_json(self):
        self.client.force_authenticate(user=self.admin_alpha)
        response = self.client.get('/api/v1/foundation/enrolment?group_by=campus&from=2025-08-01&to=2025-08-31')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.json()
        self.assertEqual(data["group_by"], "campus")
        self.assertIn("summary", data)
        self.assertIn("data", data)
        self.assertEqual(data["summary"]["total_prospects"], 1)
        self.assertEqual(data["summary"]["total_accepted"], 1)

    def test_endpoint_invalid_group_by(self):
        self.client.force_authenticate(user=self.admin_alpha)
        response = self.client.get('/api/v1/foundation/enrolment?group_by=invalid_param')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("group_by", response.json().get("detail", ""))

    def test_endpoint_invalid_date_format(self):
        self.client.force_authenticate(user=self.admin_alpha)
        response = self.client.get('/api/v1/foundation/enrolment?from=not-a-date')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("from", response.json().get("detail", ""))

    def test_endpoint_csv_export_fnd014_header(self):
        self.client.force_authenticate(user=self.admin_alpha)
        response = self.client.get('/api/v1/foundation/enrolment?format=csv&group_by=campus')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("text/csv", response["Content-Type"])
        self.assertIn("attachment; filename=", response["Content-Disposition"])
        content = response.content.decode("utf-8")
        # Check FND-014 audit comments
        self.assertIn("Laporan Pipa Pendaftaran", content)
        self.assertIn("# Waktu Dibuat :", content)
        self.assertIn(f"# Pengguna     : {self.admin_alpha.full_name}", content)
        self.assertIn("# Filter       :", content)
        # Check CSV header row
        self.assertIn("Nama Sekolah,NPSN,Jenjang,Tingkat Kelas,Calon Siswa (Prospects),Diterima (Accepted),Siswa Aktif,Keluar (Churned),Tingkat Konversi (%),Tingkat Retensi (%)", content)

    def test_endpoint_xlsx_export(self):
        self.client.force_authenticate(user=self.admin_alpha)
        response = self.client.get('/api/v1/foundation/enrolment?format=xlsx&group_by=grade')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", response["Content-Type"])
        self.assertIn("attachment; filename=", response["Content-Disposition"])
        self.assertTrue(len(response.content) > 0)

        # Inspect openpyxl workbook
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.cell(row=1, column=1).value, "Laporan Pipa Pendaftaran Yayasan (EduCore)")
        self.assertTrue(ws.cell(row=2, column=1).value.startswith("Waktu Dibuat :"))
        self.assertTrue(ws.cell(row=3, column=1).value.startswith("Pengguna     :"))

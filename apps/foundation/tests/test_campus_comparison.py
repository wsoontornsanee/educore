"""Tests for Foundation Campus Comparison endpoint (spec/03 §2, §5, FND-004, FND-005, FND-006, FND-014)."""
import csv
from datetime import date, timedelta
from decimal import Decimal
import io
from django.utils import timezone
import openpyxl
from rest_framework import status
from rest_framework.test import APITestCase
from apps.foundation.models import RptFoundationKPI
from apps.identity.models import Foundation, School, User
from apps.identity.rbac import (
    assign_role, ROLE_FOUNDATION_ADMIN, ROLE_PARENT, SCOPE_FOUNDATION
)
from educore.middleware.tenancy import clear_current_foundation_id, tenant_context

class CampusComparisonTests(APITestCase):
    def setUp(self):
        clear_current_foundation_id()
        # 1. Foundation Alpha
        self.foundation_alpha = Foundation.objects.create(
            legal_name="Yayasan Alpha Nusantara",
            brand_name="Alpha Nusantara",
            npwp="01.111.222.3-001.000",
            reporting_currency="IDR",
        )
        self.school_a1 = School.all_tenants.create(
            foundation_id=self.foundation_alpha.id,
            name="SD Alpha Satu",
            npsn="10100001",
            level=School.LEVEL_SD,
            base_currency="IDR",
            is_active=True,
        )
        self.school_a2 = School.all_tenants.create(
            foundation_id=self.foundation_alpha.id,
            name="SMP Alpha Dua",
            npsn="10100002",
            level=School.LEVEL_SMP,
            base_currency="IDR",
            is_active=True,
        )
        self.school_a3 = School.all_tenants.create(
            foundation_id=self.foundation_alpha.id,
            name="SMA Alpha Tiga",
            npsn="10100003",
            level=School.LEVEL_SMA,
            base_currency="IDR",
            is_active=True,
        )

        # 2. Foundation Beta (for cross-tenant testing)
        self.foundation_beta = Foundation.objects.create(
            legal_name="Yayasan Beta Cendekia",
            brand_name="Beta Cendekia",
            reporting_currency="IDR",
        )
        self.school_b1 = School.all_tenants.create(
            foundation_id=self.foundation_beta.id,
            name="SMA Beta Cendekia",
            npsn="20200001",
            level=School.LEVEL_SMA,
            base_currency="IDR",
            is_active=True,
        )

        # 3. Users
        self.admin_alpha = User.all_tenants.create_user(
            phone_e164="+6281111111111",
            foundation_id=self.foundation_alpha.id,
            full_name="Ketua Yayasan Alpha",
        )
        assign_role(
            user=self.admin_alpha,
            role=ROLE_FOUNDATION_ADMIN,
            scope_type=SCOPE_FOUNDATION,
            scope_id=self.foundation_alpha.id,
            foundation_id=self.foundation_alpha.id,
        )

        self.parent_user = User.all_tenants.create_user(
            phone_e164="+6281222222222",
            foundation_id=self.foundation_alpha.id,
            full_name="Wali Murid Alpha",
        )
        assign_role(
            user=self.parent_user,
            role=ROLE_PARENT,
            scope_type=SCOPE_FOUNDATION,
            scope_id=self.foundation_alpha.id,
            foundation_id=self.foundation_alpha.id,
        )

        # 4. Seed Rollup KPIs for September 2026
        # School A1: 90% collection rate
        self.kpi_a1 = RptFoundationKPI.objects.create(
            foundation_id=self.foundation_alpha.id,
            school_id=self.school_a1.id,
            period_start=date(2026, 9, 1),
            period_end=date(2026, 9, 30),
            billed=Decimal('100000000.00'),
            collected=Decimal('90000000.00'),
            outstanding=Decimal('10000000.00'),
            ar_0_30=Decimal('8000000.00'),
            ar_31_60=Decimal('2000000.00'),
            campus_spend=Decimal('15000000.00'),
            active_students=500,
            avg_attendance_pct=Decimal('98.00'),
            currency='IDR',
            reporting_currency='IDR',
        )

        # School A2: 80% collection rate
        self.kpi_a2 = RptFoundationKPI.objects.create(
            foundation_id=self.foundation_alpha.id,
            school_id=self.school_a2.id,
            period_start=date(2026, 9, 1),
            period_end=date(2026, 9, 30),
            billed=Decimal('50000000.00'),
            collected=Decimal('40000000.00'),
            outstanding=Decimal('10000000.00'),
            ar_0_30=Decimal('5000000.00'),
            ar_31_60=Decimal('5000000.00'),
            campus_spend=Decimal('8000000.00'),
            active_students=300,
            avg_attendance_pct=Decimal('95.00'),
            currency='IDR',
            reporting_currency='IDR',
        )

        # School A3: 70% collection rate
        self.kpi_a3 = RptFoundationKPI.objects.create(
            foundation_id=self.foundation_alpha.id,
            school_id=self.school_a3.id,
            period_start=date(2026, 9, 1),
            period_end=date(2026, 9, 30),
            billed=Decimal('20000000.00'),
            collected=Decimal('14000000.00'),
            outstanding=Decimal('6000000.00'),
            ar_0_30=Decimal('4000000.00'),
            ar_31_60=Decimal('2000000.00'),
            campus_spend=Decimal('3000000.00'),
            active_students=150,
            avg_attendance_pct=Decimal('92.00'),
            currency='IDR',
            reporting_currency='IDR',
        )

        # Seed Beta KPI (cross-tenant check)
        self.kpi_b1 = RptFoundationKPI.objects.create(
            foundation_id=self.foundation_beta.id,
            school_id=self.school_b1.id,
            period_start=date(2026, 9, 1),
            period_end=date(2026, 9, 30),
            billed=Decimal('80000000.00'),
            collected=Decimal('76000000.00'),
            outstanding=Decimal('4000000.00'),
            active_students=400,
            avg_attendance_pct=Decimal('96.00'),
            currency='IDR',
            reporting_currency='IDR',
        )

    def tearDown(self):
        clear_current_foundation_id()

    def test_unauthenticated_request_denied(self):
        """Unauthenticated requests must receive 401 Unauthorized."""
        response = self.client.get('/api/v1/foundation/schools/compare')
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_unauthorized_role_denied(self):
        """Users without school_config.read (e.g. parent) must receive 403 Forbidden."""
        self.client.force_authenticate(user=self.parent_user)
        with tenant_context(self.foundation_alpha.id):
            response = self.client.get('/api/v1/foundation/schools/compare')
            self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_cross_tenant_isolation(self):
        """Alpha foundation admin must never see Beta schools or Beta KPI data."""
        self.client.force_authenticate(user=self.admin_alpha)
        with tenant_context(self.foundation_alpha.id):
            response = self.client.get('/api/v1/foundation/schools/compare')
            self.assertEqual(response.status_code, status.HTTP_200_OK)
            results = response.data['results']
            school_ids = [r['school_id'] for r in results]
            self.assertIn(self.school_a1.id, school_ids)
            self.assertIn(self.school_a2.id, school_ids)
            self.assertIn(self.school_a3.id, school_ids)
            self.assertNotIn(self.school_b1.id, school_ids)

            # Explicitly attempting to query Beta school_id must yield empty results
            res_cross = self.client.get(f'/api/v1/foundation/schools/compare?school_ids={self.school_b1.id}')
            self.assertEqual(res_cross.status_code, status.HTTP_200_OK)
            self.assertEqual(len(res_cross.data['results']), 0)

    def test_default_sorting_by_collection_rate_desc(self):
        """Default ranking is collection_rate descending (rank 1 = highest collection rate)."""
        self.client.force_authenticate(user=self.admin_alpha)
        with tenant_context(self.foundation_alpha.id):
            response = self.client.get('/api/v1/foundation/schools/compare')
            self.assertEqual(response.status_code, status.HTTP_200_OK)
            results = response.data['results']
            self.assertEqual(len(results), 3)

            # A1 (90%) -> rank 1, A2 (80%) -> rank 2, A3 (70%) -> rank 3
            self.assertEqual(results[0]['school_id'], self.school_a1.id)
            self.assertEqual(results[0]['rank'], 1)
            self.assertEqual(results[0]['collection_rate_pct'], '90.00')

            self.assertEqual(results[1]['school_id'], self.school_a2.id)
            self.assertEqual(results[1]['rank'], 2)
            self.assertEqual(results[1]['collection_rate_pct'], '80.00')

            self.assertEqual(results[2]['school_id'], self.school_a3.id)
            self.assertEqual(results[2]['rank'], 3)
            self.assertEqual(results[2]['collection_rate_pct'], '70.00')

            # Summary checks
            summary = response.data['summary']
            self.assertEqual(summary['total_schools'], 3)
            self.assertEqual(summary['billed'], '170000000.00')
            self.assertEqual(summary['collected'], '144000000.00')
            # 144 / 170 * 100 = 84.71%
            self.assertEqual(summary['collection_rate_pct'], '84.71')
            self.assertEqual(summary['outstanding'], '26000000.00')
            self.assertEqual(summary['campus_spend'], '26000000.00')
            self.assertEqual(summary['active_students'], 950)
            self.assertEqual(summary['reporting_currency'], 'IDR')

    def test_sorting_by_other_metrics(self):
        """Campus comparison can sort by any supported metric (billed, active_students, etc.)."""
        self.client.force_authenticate(user=self.admin_alpha)
        with tenant_context(self.foundation_alpha.id):
            # 1. Sort by billed descending
            res_billed = self.client.get('/api/v1/foundation/schools/compare?metric=billed&order=desc')
            self.assertEqual(res_billed.status_code, status.HTTP_200_OK)
            self.assertEqual(res_billed.data['results'][0]['school_id'], self.school_a1.id)
            self.assertEqual(res_billed.data['results'][0]['billed'], '100000000.00')

            # 2. Sort by active_students descending
            res_students = self.client.get('/api/v1/foundation/schools/compare?metric=active_students')
            self.assertEqual(res_students.status_code, status.HTTP_200_OK)
            self.assertEqual(res_students.data['results'][0]['active_students'], 500)
            self.assertEqual(res_students.data['results'][2]['active_students'], 150)

            # 3. Sort by collection_rate ascending (order=asc)
            res_asc = self.client.get('/api/v1/foundation/schools/compare?metric=collection_rate&order=asc')
            self.assertEqual(res_asc.status_code, status.HTTP_200_OK)
            self.assertEqual(res_asc.data['results'][0]['school_id'], self.school_a3.id)
            self.assertEqual(res_asc.data['results'][0]['collection_rate_pct'], '70.00')

            # 4. Sort by name ascending
            res_name = self.client.get('/api/v1/foundation/schools/compare?metric=name')
            self.assertEqual(res_name.status_code, status.HTTP_200_OK)
            names = [r['school_name'] for r in res_name.data['results']]
            self.assertEqual(names, sorted(names))

    def test_date_range_filtering(self):
        """Filter by from and to limits rollup aggregation to target window."""
        # Seed an August row for A1
        RptFoundationKPI.objects.create(
            foundation_id=self.foundation_alpha.id,
            school_id=self.school_a1.id,
            period_start=date(2026, 8, 1),
            period_end=date(2026, 8, 31),
            billed=Decimal('50000000.00'),
            collected=Decimal('50000000.00'),
            outstanding=Decimal('0.00'),
            active_students=490,
            avg_attendance_pct=Decimal('97.00'),
            currency='IDR',
            reporting_currency='IDR',
        )

        self.client.force_authenticate(user=self.admin_alpha)
        with tenant_context(self.foundation_alpha.id):
            # Target August only
            res_aug = self.client.get('/api/v1/foundation/schools/compare?from=2026-08-01&to=2026-08-31')
            self.assertEqual(res_aug.status_code, status.HTTP_200_OK)
            a1_row = next(r for r in res_aug.data['results'] if r['school_id'] == self.school_a1.id)
            self.assertEqual(a1_row['billed'], '50000000.00')
            self.assertEqual(a1_row['collected'], '50000000.00')

            # Target both months
            res_both = self.client.get('/api/v1/foundation/schools/compare?from=2026-08-01&to=2026-09-30')
            a1_both = next(r for r in res_both.data['results'] if r['school_id'] == self.school_a1.id)
            self.assertEqual(a1_both['billed'], '150000000.00')
            self.assertEqual(a1_both['collected'], '140000000.00')

    def test_school_ids_filtering(self):
        """Filtering by school_ids restricts comparison to requested subset."""
        self.client.force_authenticate(user=self.admin_alpha)
        with tenant_context(self.foundation_alpha.id):
            response = self.client.get(
                f'/api/v1/foundation/schools/compare?school_ids={self.school_a1.id},{self.school_a3.id}'
            )
            self.assertEqual(response.status_code, status.HTTP_200_OK)
            results = response.data['results']
            self.assertEqual(len(results), 2)
            returned_ids = {r['school_id'] for r in results}
            self.assertEqual(returned_ids, {self.school_a1.id, self.school_a3.id})

    def test_school_without_kpi_record_included_with_zeroes(self):
        """A school with no rollup records yet must still appear with 0.00 metrics."""
        school_new = School.all_tenants.create(
            foundation_id=self.foundation_alpha.id,
            name="SMK Alpha Baru",
            npsn="10100004",
            level=School.LEVEL_SMK,
            is_active=True,
        )
        self.client.force_authenticate(user=self.admin_alpha)
        with tenant_context(self.foundation_alpha.id):
            response = self.client.get('/api/v1/foundation/schools/compare')
            self.assertEqual(response.status_code, status.HTTP_200_OK)
            results = response.data['results']
            self.assertEqual(len(results), 4)
            new_row = next(r for r in results if r['school_id'] == school_new.id)
            self.assertEqual(new_row['billed'], '0.00')
            self.assertEqual(new_row['collected'], '0.00')
            self.assertEqual(new_row['collection_rate_pct'], '0.00')
            self.assertEqual(new_row['active_students'], 0)
            self.assertEqual(new_row['rank'], 4)

    def test_freshness_evaluation_and_stale_flag(self):
        """FND-006: Freshness timestamp surfaced and stale flag raised beyond 15 minutes."""
        self.client.force_authenticate(user=self.admin_alpha)
        with tenant_context(self.foundation_alpha.id):
            response = self.client.get('/api/v1/foundation/schools/compare')
            freshness = response.data['freshness']
            self.assertIsNotNone(freshness['computed_at'])
            self.assertFalse(freshness['stale'])

            # Stale simulation
            RptFoundationKPI.objects.filter(foundation_id=self.foundation_alpha.id).update(
                computed_at=timezone.now() - timedelta(minutes=20)
            )
            response_stale = self.client.get('/api/v1/foundation/schools/compare')
            self.assertTrue(response_stale.data['freshness']['stale'])

    def test_25_schools_capacity_fnd_004(self):
        """FND-004: Campus comparison must support up to 25 schools without pagination lag."""
        # Create 22 more schools to reach 25 total
        created_schools = []
        for i in range(4, 26):
            s = School.all_tenants.create(
                foundation_id=self.foundation_alpha.id,
                name=f"Sekolah Kampus {i:02d}",
                npsn=f"1010{i:04d}",
                level=School.LEVEL_SD,
                is_active=True,
            )
            created_schools.append(s)
            RptFoundationKPI.objects.create(
                foundation_id=self.foundation_alpha.id,
                school_id=s.id,
                period_start=date(2026, 9, 1),
                period_end=date(2026, 9, 30),
                billed=Decimal(f"{i * 1000000}.00"),
                collected=Decimal(f"{i * 800000}.00"),
                outstanding=Decimal(f"{i * 200000}.00"),
                active_students=i * 20,
                avg_attendance_pct=Decimal("95.00"),
                currency='IDR',
                reporting_currency='IDR',
            )

        self.client.force_authenticate(user=self.admin_alpha)
        with tenant_context(self.foundation_alpha.id):
            response = self.client.get('/api/v1/foundation/schools/compare')
            self.assertEqual(response.status_code, status.HTTP_200_OK)
            results = response.data['results']
            self.assertEqual(len(results), 25)
            # Verify ranks 1 through 25 are consecutively assigned
            ranks = [r['rank'] for r in results]
            self.assertEqual(ranks, list(range(1, 26)))

    def test_invalid_parameters_return_400(self):
        """Invalid date formats, metric names, or order parameters return 400 Bad Request."""
        self.client.force_authenticate(user=self.admin_alpha)
        with tenant_context(self.foundation_alpha.id):
            # Invalid from date
            res1 = self.client.get('/api/v1/foundation/schools/compare?from=bad-date')
            self.assertEqual(res1.status_code, status.HTTP_400_BAD_REQUEST)

            # from > to
            res2 = self.client.get('/api/v1/foundation/schools/compare?from=2026-09-30&to=2026-09-01')
            self.assertEqual(res2.status_code, status.HTTP_400_BAD_REQUEST)

            # Invalid metric
            res3 = self.client.get('/api/v1/foundation/schools/compare?metric=unknown_kpi')
            self.assertEqual(res3.status_code, status.HTTP_400_BAD_REQUEST)

            # Invalid order
            res4 = self.client.get('/api/v1/foundation/schools/compare?order=sideways')
            self.assertEqual(res4.status_code, status.HTTP_400_BAD_REQUEST)

    def test_csv_export_format_fnd_014(self):
        """format=csv streams RFC 4180 CSV with mandatory FND-014 audit headers."""
        self.client.force_authenticate(user=self.admin_alpha)
        with tenant_context(self.foundation_alpha.id):
            response = self.client.get('/api/v1/foundation/schools/compare?format=csv&metric=collection_rate')
            self.assertEqual(response.status_code, status.HTTP_200_OK)
            self.assertIn('text/csv', response['Content-Type'])
            self.assertIn('attachment; filename="perbandingan_kampus.csv"', response['Content-Disposition'])

            content = response.content.decode('utf-8-sig')
            # FND-014: Check audit headers
            self.assertIn("# Laporan Perbandingan Kampus Yayasan", content)
            self.assertIn("# Waktu Dibuat:", content)
            self.assertIn("# Pengguna: Ketua Yayasan Alpha", content)
            self.assertIn("# Filter: Dari=Semua, Sampai=Semua, Metrik=collection_rate", content)

            # Parse tabular part
            table_lines = [line for line in content.splitlines() if not line.startswith('#') and line.strip()]
            reader = csv.reader(table_lines)
            rows = list(reader)
            # Header + 3 school rows + 1 summary footer row = 5 rows
            self.assertEqual(len(rows), 5)
            self.assertEqual(rows[0][0], "Peringkat")
            self.assertEqual(rows[1][1], self.school_a1.name)
            self.assertEqual(rows[1][7], "90.00")
            self.assertEqual(rows[4][0], "Total / Rata-rata")

    def test_xlsx_export_format_fnd_014(self):
        """format=xlsx streams Excel workbook with styled audit banner, table, and formulas."""
        self.client.force_authenticate(user=self.admin_alpha)
        with tenant_context(self.foundation_alpha.id):
            response = self.client.get('/api/v1/foundation/schools/compare?format=xlsx&metric=collection_rate')
            self.assertEqual(response.status_code, status.HTTP_200_OK)
            self.assertIn(
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                response['Content-Type']
            )
            self.assertIn('attachment; filename="perbandingan_kampus.xlsx"', response['Content-Disposition'])

            # Verify with openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(response.content))
            ws = wb.active
            self.assertEqual(ws.title, "Perbandingan Kampus")

            # FND-014: Metadata header banner
            self.assertIn("Laporan Perbandingan Kampus", ws.cell(row=1, column=1).value)
            self.assertIn("Ketua Yayasan Alpha", ws.cell(row=3, column=1).value)
            self.assertIn("Metrik=collection_rate", ws.cell(row=4, column=1).value)

            # Table Header at row 6
            self.assertEqual(ws.cell(row=6, column=1).value, "Peringkat")
            self.assertEqual(ws.cell(row=6, column=2).value, "Nama Sekolah")

            # First data row at row 7
            self.assertEqual(ws.cell(row=7, column=1).value, 1)
            self.assertEqual(ws.cell(row=7, column=2).value, self.school_a1.name)

            # Summary row at row 10
            self.assertEqual(ws.cell(row=10, column=1).value, "Total / Rata-rata")

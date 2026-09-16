"""Tests for Foundation Dashboard PDF/XLSX exports (FND-014, RPT-002/003)."""
import io
from datetime import date
from decimal import Decimal
from unittest import mock

from django.core.management import call_command
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from apps.core.models import ExportJob
from apps.foundation.models import RptFoundationKPI
from apps.identity.models import Foundation, School, User
from apps.identity.rbac import assign_role, ROLE_FOUNDATION_ADMIN, SCOPE_FOUNDATION
from apps.notifications.models import NotificationIntent
from educore.middleware.tenancy import clear_current_foundation_id, tenant_context


class FoundationDashboardExportTests(APITestCase):
    def setUp(self):
        clear_current_foundation_id()
        self.foundation = Foundation.objects.create(
            legal_name="Yayasan Bina Bangsa", brand_name="Bina Bangsa", npwp="01.888.777.6-001.000",
        )
        self.other_foundation = Foundation.objects.create(
            legal_name="Yayasan Lain", brand_name="Lain", npwp="02.111.222.3-001.000",
        )
        self.school = School.all_tenants.create(
            foundation_id=self.foundation.id, name="SMA Bina Bangsa", npsn="40100001", level=School.LEVEL_SMA,
        )
        self.admin = User.all_tenants.create_user(
            phone_e164="+6281999999999", foundation_id=self.foundation.id,
            full_name="Ketua Yayasan Bina Bangsa",
        )
        assign_role(
            user=self.admin, role=ROLE_FOUNDATION_ADMIN, scope_type=SCOPE_FOUNDATION,
            scope_id=self.foundation.id, foundation_id=self.foundation.id,
        )
        RptFoundationKPI.objects.create(
            foundation_id=self.foundation.id, school_id=self.school.id,
            period_start=date(2026, 9, 1), period_end=date(2026, 9, 30),
            billed=Decimal('50000000.00'), collected=Decimal('42500000.00'),
            outstanding=Decimal('7500000.00'), active_students=450,
            avg_attendance_pct=Decimal('96.50'), currency='IDR', reporting_currency='IDR',
        )

    def tearDown(self):
        clear_current_foundation_id()

    def test_post_export_requires_valid_format(self):
        self.client.force_authenticate(user=self.admin)
        with tenant_context(self.foundation.id):
            response = self.client.post('/api/v1/foundation/exports', data={'format': 'DOCX'}, format='json')
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_post_export_requires_registered_report(self):
        self.client.force_authenticate(user=self.admin)
        with tenant_context(self.foundation.id):
            response = self.client.post(
                '/api/v1/foundation/exports', data={'format': 'xlsx', 'report': 'not_a_real_report'}, format='json',
            )
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_post_export_rejects_non_string_report(self):
        self.client.force_authenticate(user=self.admin)
        with tenant_context(self.foundation.id):
            response = self.client.post(
                '/api/v1/foundation/exports', data={'format': 'xlsx', 'report': {'x': 1}}, format='json',
            )
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_post_export_rejects_format_not_supported_by_report(self):
        """A format that's globally valid but not registered for this report_key
        (e.g. CSV for the dashboard, which only renders PDF/XLSX) must 400,
        not silently render a mismatched file (FND-010 review finding)."""
        self.client.force_authenticate(user=self.admin)
        with tenant_context(self.foundation.id):
            response = self.client.post(
                '/api/v1/foundation/exports',
                data={'format': 'csv', 'report': 'foundation_dashboard'}, format='json',
            )
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_post_export_enqueues_job(self):
        self.client.force_authenticate(user=self.admin)
        with tenant_context(self.foundation.id):
            response = self.client.post(
                '/api/v1/foundation/exports',
                data={'format': 'xlsx', 'filters': {'from': '2026-09-01', 'to': '2026-09-30'}},
                format='json',
            )
            self.assertEqual(response.status_code, status.HTTP_202_ACCEPTED)
            job_id = response.data['job_id']

        job = ExportJob.all_tenants.get(id=job_id)
        self.assertEqual(job.foundation_id, self.foundation.id)
        self.assertEqual(job.report_key, 'foundation_dashboard')
        self.assertEqual(job.format, ExportJob.FORMAT_XLSX)
        self.assertEqual(job.requested_by, str(self.admin.id))
        self.assertEqual(job.requested_by_name, self.admin.full_name)

    def test_get_export_status_before_completion(self):
        self.client.force_authenticate(user=self.admin)
        with tenant_context(self.foundation.id):
            post_response = self.client.post('/api/v1/foundation/exports', data={'format': 'PDF'}, format='json')
            job_id = post_response.data['job_id']

            response = self.client.get(f'/api/v1/foundation/exports/{job_id}')
            self.assertEqual(response.status_code, status.HTTP_200_OK)
            self.assertEqual(response.data['status'], ExportJob.STATUS_PENDING)
            self.assertIsNone(response.data['download_url'])

    def test_get_export_status_cross_tenant_returns_404(self):
        other_admin = User.all_tenants.create_user(
            phone_e164="+6281888888888", foundation_id=self.other_foundation.id, full_name="Admin Lain",
        )
        assign_role(
            user=other_admin, role=ROLE_FOUNDATION_ADMIN, scope_type=SCOPE_FOUNDATION,
            scope_id=self.other_foundation.id, foundation_id=self.other_foundation.id,
        )
        with tenant_context(self.foundation.id):
            job = ExportJob.objects.create(
                foundation_id=self.foundation.id, report_key='foundation_dashboard', format=ExportJob.FORMAT_PDF,
            )

        self.client.force_authenticate(user=other_admin)
        with tenant_context(self.other_foundation.id):
            response = self.client.get(f'/api/v1/foundation/exports/{job.id}')
            self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    @mock.patch('apps.core.storage._client')
    def test_full_pipeline_xlsx_export_completes_and_notifies(self, mock_client):
        mock_client.return_value.bucket.return_value.blob.return_value.generate_signed_url.return_value = (
            'https://signed.example/download.xlsx'
        )
        self.client.force_authenticate(user=self.admin)
        with tenant_context(self.foundation.id):
            post_response = self.client.post(
                '/api/v1/foundation/exports',
                data={'format': 'xlsx', 'filters': {'school_ids': [self.school.id]}},
                format='json',
            )
            job_id = post_response.data['job_id']

        call_command('drain_tasks', limit=10)

        job = ExportJob.all_tenants.get(id=job_id)
        self.assertEqual(job.status, ExportJob.STATUS_COMPLETED, job.error_text)
        self.assertTrue(job.result_key)

        with tenant_context(self.foundation.id):
            response = self.client.get(f'/api/v1/foundation/exports/{job_id}')
            self.assertEqual(response.data['status'], ExportJob.STATUS_COMPLETED)
            self.assertIsNotNone(response.data['download_url'])

        intent = NotificationIntent.all_tenants.filter(
            foundation_id=self.foundation.id, template_key='core.export.ready',
        ).first()
        self.assertIsNotNone(intent)
        self.assertEqual(intent.recipient_user_id, self.admin.id)

    @mock.patch('apps.core.storage._client')
    def test_xlsx_export_defuses_formula_injection_in_filters(self, mock_client):
        """A filters value starting with =/+/-/@ must not be interpretable as a
        formula by the spreadsheet app that opens the exported file."""
        mock_client.return_value.bucket.return_value.blob.return_value.generate_signed_url.return_value = (
            'https://signed.example/download.xlsx'
        )
        self.client.force_authenticate(user=self.admin)
        with tenant_context(self.foundation.id):
            post_response = self.client.post(
                '/api/v1/foundation/exports',
                data={'format': 'xlsx', 'filters': {'=cmd|/c calc': 'x'}},
                format='json',
            )
            job_id = post_response.data['job_id']

        call_command('drain_tasks', limit=10)

        job = ExportJob.all_tenants.get(id=job_id)
        self.assertEqual(job.status, ExportJob.STATUS_COMPLETED, job.error_text)

        upload_call = mock_client.return_value.bucket.return_value.blob.return_value.upload_from_string
        data = upload_call.call_args[0][0]
        wb = load_workbook(filename=io.BytesIO(data))
        ws = wb.active
        filters_cell = next(row[1] for row in ws.iter_rows() if row[0].value == 'Filter')
        self.assertTrue(str(filters_cell.value).startswith("'="))

    @mock.patch('apps.core.storage._client')
    def test_full_pipeline_pdf_export_completes(self, mock_client):
        mock_client.return_value.bucket.return_value.blob.return_value.generate_signed_url.return_value = (
            'https://signed.example/download.pdf'
        )
        self.client.force_authenticate(user=self.admin)
        with tenant_context(self.foundation.id):
            post_response = self.client.post('/api/v1/foundation/exports', data={'format': 'PDF'}, format='json')
            job_id = post_response.data['job_id']

        call_command('drain_tasks', limit=10)

        job = ExportJob.all_tenants.get(id=job_id)
        self.assertEqual(job.status, ExportJob.STATUS_COMPLETED, job.error_text)
        self.assertIn(job.result_key.rsplit('.', 1)[-1], ('pdf', 'html'))

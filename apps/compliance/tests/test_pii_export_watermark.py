"""Automated tests for Universal PII Export Watermarking & Access Log (spec/14 §3 CMP-016, spec/15 §2 RPT-004)."""
from datetime import date, timedelta
import io
from unittest import mock

from django.test import TestCase
from django.utils import timezone
from rest_framework import status
from rest_framework.test import APIClient
import openpyxl

from apps.compliance.exports import REPORT_KEY_DAPODIK, REPORT_KEY_EMIS
from apps.compliance.models import PiiExportAccessLog, StatutorySystem
from apps.compliance.tests.test_statutory_export import StatutoryExportFixtures
from apps.core.models import AuditEvent, ExportJob
from apps.core.services import (
    create_export_job,
    is_export_pii,
    register_export_renderer,
    run_export_job,
)
from apps.identity.models import Foundation, User
from apps.identity.rbac import ROLE_FOUNDATION_ADMIN, ROLE_TEACHER, SCOPE_FOUNDATION, assign_role
from educore.middleware.tenancy import clear_current_foundation_id, tenant_context


class StatutoryPiiRegistrationTests(TestCase):
    def test_dapodik_and_emis_registered_as_pii(self):
        self.assertTrue(is_export_pii(REPORT_KEY_DAPODIK))
        self.assertTrue(is_export_pii(REPORT_KEY_EMIS))

    def test_unregistered_reports_not_pii(self):
        self.assertFalse(is_export_pii('foundation_dashboard'))
        self.assertFalse(is_export_pii('random_report_key'))


class PiiExportWatermarkTests(StatutoryExportFixtures):
    def setUp(self):
        super().setUp()
        self.admin_user = User.all_tenants.create_user(
            phone_e164="+6281122334455",
            foundation_id=self.foundation.id,
            full_name="Admin Tata Usaha",
        )

    @mock.patch('apps.core.storage.upload_bytes')
    def test_dapodik_xlsx_watermarked(self, mock_upload):
        uploaded_payloads = []

        def capture_upload(key, data, content_type):
            uploaded_payloads.append((key, data, content_type))

        mock_upload.side_effect = capture_upload

        job = create_export_job(
            report_key=REPORT_KEY_DAPODIK,
            export_format=ExportJob.FORMAT_XLSX,
            filters={'school_id': self.school.id, 'system': StatutorySystem.DAPODIK},
            foundation_id=self.foundation.id,
            requested_by=str(self.admin_user.id),
            requested_by_name=self.admin_user.full_name,
        )

        run_export_job({'export_job_id': job.id})
        job.refresh_from_db()
        self.assertEqual(job.status, ExportJob.STATUS_COMPLETED)
        self.assertTrue(job.result_key)
        self.assertEqual(len(uploaded_payloads), 1)

        _key, data, content_type = uploaded_payloads[0]
        self.assertIn('spreadsheetml', content_type)

        # Inspect XLSX content via openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data))

        # 1. Document properties
        self.assertEqual(wb.properties.creator, "Admin Tata Usaha")
        self.assertIn("RAHASIA PII", wb.properties.description)
        self.assertIn(str(self.admin_user.id), wb.properties.description)

        # 2. Info sheet watermark block
        self.assertIn('Info', wb.sheetnames)
        info_ws = wb['Info']
        info_text = " ".join(str(cell.value) for row in info_ws.iter_rows() for cell in row if cell.value)
        self.assertIn("RAHASIA / PII", info_text)
        self.assertIn("Admin Tata Usaha", info_text)
        self.assertIn(f"ExportJob#{job.id}", info_text)

        # 3. Worksheets print headers/footers
        for ws in wb.worksheets:
            self.assertIn("RAHASIA PII", ws.oddHeader.left.text)
            self.assertIn(f"ExportJob#{job.id}", ws.oddFooter.center.text)

    @mock.patch('apps.core.storage.upload_bytes')
    def test_dapodik_csv_watermarked(self, mock_upload):
        uploaded_payloads = []

        def capture_upload(key, data, content_type):
            uploaded_payloads.append((key, data, content_type))

        mock_upload.side_effect = capture_upload

        job = create_export_job(
            report_key=REPORT_KEY_DAPODIK,
            export_format=ExportJob.FORMAT_CSV,
            filters={'school_id': self.school.id, 'system': StatutorySystem.DAPODIK, 'sheet': 'students'},
            foundation_id=self.foundation.id,
            requested_by=str(self.admin_user.id),
            requested_by_name=self.admin_user.full_name,
        )

        run_export_job({'export_job_id': job.id})
        job.refresh_from_db()
        self.assertEqual(job.status, ExportJob.STATUS_COMPLETED)

        _key, data, content_type = uploaded_payloads[0]
        self.assertEqual(content_type, 'text/csv')

        csv_text = data.decode('utf-8-sig')

        # Line 1 columns preserved
        first_line = csv_text.strip().splitlines()[0]
        self.assertIn('nisn', first_line)
        self.assertIn('nama', first_line)

        # Student row preserved
        self.assertIn('Andi Wijaya', csv_text)

        # Watermark banner present
        self.assertIn("RAHASIA PII — UU No. 27/2022 (UU PDP)", csv_text)
        self.assertIn("Admin Tata Usaha", csv_text)
        self.assertIn(f"ExportJob#{job.id}", csv_text)

    @mock.patch('apps.core.storage.upload_bytes')
    def test_pii_export_access_log_recorded(self, mock_upload):
        job = create_export_job(
            report_key=REPORT_KEY_DAPODIK,
            export_format=ExportJob.FORMAT_XLSX,
            filters={'school_id': self.school.id, 'system': StatutorySystem.DAPODIK},
            foundation_id=self.foundation.id,
            requested_by=str(self.admin_user.id),
            requested_by_name=self.admin_user.full_name,
        )

        run_export_job({'export_job_id': job.id})

        log = PiiExportAccessLog.objects.filter(export_job=job).first()
        self.assertIsNotNone(log)
        self.assertEqual(log.foundation_id, self.foundation.id)
        self.assertEqual(log.report_key, REPORT_KEY_DAPODIK)
        self.assertEqual(log.format, ExportJob.FORMAT_XLSX)
        self.assertEqual(log.exported_by_id, str(self.admin_user.id))
        self.assertEqual(log.exported_by_name, self.admin_user.full_name)
        self.assertEqual(log.school_id, self.school.id)
        self.assertGreater(log.record_count, 0)
        self.assertGreater(log.file_size, 0)
        self.assertIn("RAHASIA PII", log.watermark_text)
        self.assertIn("Admin Tata Usaha", log.watermark_text)

        # Verify AuditEvent was logged
        audit_event = AuditEvent.objects.filter(
            action='EXPORT_PII',
            entity_type='ExportJob',
            entity_id=str(job.id),
        ).first()
        self.assertIsNotNone(audit_event)
        self.assertEqual(audit_event.foundation_id, self.foundation.id)
        self.assertEqual(audit_event.diff['report_key'], REPORT_KEY_DAPODIK)
        self.assertEqual(sorted(audit_event.diff['pii_types']), ['NIK', 'NISN'])

    @mock.patch('apps.core.storage.upload_bytes')
    def test_no_nik_in_access_log_or_audit(self, mock_upload):
        """spec/14 §7 criterion 6: No log or audit payload contains a NIK."""
        job = create_export_job(
            report_key=REPORT_KEY_DAPODIK,
            export_format=ExportJob.FORMAT_XLSX,
            filters={'school_id': self.school.id, 'system': StatutorySystem.DAPODIK},
            foundation_id=self.foundation.id,
            requested_by=str(self.admin_user.id),
            requested_by_name=self.admin_user.full_name,
        )

        run_export_job({'export_job_id': job.id})

        log = PiiExportAccessLog.objects.get(export_job=job)
        student_nik = self.student_person.nik
        self.assertTrue(student_nik)

        # Assert NIK not in log fields
        self.assertNotIn(student_nik, log.watermark_text)
        self.assertNotIn(student_nik, str(log.filters))

        audit_event = AuditEvent.objects.get(action='EXPORT_PII', entity_id=str(job.id))
        self.assertNotIn(student_nik, str(audit_event.diff))

    @mock.patch('apps.core.storage.upload_bytes')
    def test_non_pii_export_does_not_create_pii_log(self, mock_upload):
        @register_export_renderer('test.plain.report')
        def plain_renderer(job):
            return b'plain content', 'text/plain', 'plain.txt'

        job = create_export_job(
            report_key='test.plain.report',
            export_format=ExportJob.FORMAT_CSV,
            filters={},
            foundation_id=self.foundation.id,
            requested_by=str(self.admin_user.id),
        )

        run_export_job({'export_job_id': job.id})
        self.assertFalse(PiiExportAccessLog.objects.filter(export_job=job).exists())


class PiiExportAccessLogApiTests(TestCase):
    def setUp(self):
        clear_current_foundation_id()
        self.client = APIClient()

        # Foundation A
        self.foundation_a = Foundation.objects.create(
            legal_name="Yayasan Harapan A",
            brand_name="Harapan A",
            npwp="01.111.111.1-001.000",
        )
        # Foundation B
        self.foundation_b = Foundation.objects.create(
            legal_name="Yayasan Mandiri B",
            brand_name="Mandiri B",
            npwp="02.222.222.2-002.000",
        )

        # Admin User A
        self.admin_a = User.all_tenants.create_user(
            phone_e164="+6281111111111",
            foundation_id=self.foundation_a.id,
            full_name="Admin Foundation A",
        )
        assign_role(
            user=self.admin_a, role=ROLE_FOUNDATION_ADMIN, scope_type=SCOPE_FOUNDATION,
            scope_id=self.foundation_a.id, foundation_id=self.foundation_a.id,
        )

        # Regular Teacher User A (no audit_log.read permission)
        self.teacher_a = User.all_tenants.create_user(
            phone_e164="+6281222222222",
            foundation_id=self.foundation_a.id,
            full_name="Guru Biasa",
        )
        assign_role(
            user=self.teacher_a, role=ROLE_TEACHER, scope_type=SCOPE_FOUNDATION,
            scope_id=self.foundation_a.id, foundation_id=self.foundation_a.id,
        )

        # Admin User B
        self.admin_b = User.all_tenants.create_user(
            phone_e164="+6281333333333",
            foundation_id=self.foundation_b.id,
            full_name="Admin Foundation B",
        )
        assign_role(
            user=self.admin_b, role=ROLE_FOUNDATION_ADMIN, scope_type=SCOPE_FOUNDATION,
            scope_id=self.foundation_b.id, foundation_id=self.foundation_b.id,
        )

        # Seed Access Logs
        with tenant_context(self.foundation_a.id):
            self.log_a1 = PiiExportAccessLog.objects.create(
                foundation_id=self.foundation_a.id,
                report_key=REPORT_KEY_DAPODIK,
                format='XLSX',
                exported_by_id=str(self.admin_a.id),
                exported_by_name=self.admin_a.full_name,
                school_id=101,
                record_count=150,
                watermark_text="RAHASIA PII A1",
                file_name="statutory_dapodik_1.xlsx",
                file_size=54321,
            )
            self.log_a2 = PiiExportAccessLog.objects.create(
                foundation_id=self.foundation_a.id,
                report_key=REPORT_KEY_EMIS,
                format='CSV',
                exported_by_id=str(self.admin_a.id),
                exported_by_name=self.admin_a.full_name,
                school_id=102,
                record_count=75,
                watermark_text="RAHASIA PII A2",
                file_name="statutory_emis_2.csv",
                file_size=12345,
            )

        with tenant_context(self.foundation_b.id):
            self.log_b1 = PiiExportAccessLog.objects.create(
                foundation_id=self.foundation_b.id,
                report_key=REPORT_KEY_DAPODIK,
                format='XLSX',
                exported_by_id=str(self.admin_b.id),
                exported_by_name=self.admin_b.full_name,
                school_id=201,
                record_count=200,
                watermark_text="RAHASIA PII B1",
                file_name="statutory_dapodik_b1.xlsx",
                file_size=88888,
            )

    def tearDown(self):
        clear_current_foundation_id()

    def test_admin_can_list_pii_export_logs(self):
        self.client.force_authenticate(user=self.admin_a)
        response = self.client.get('/api/v1/foundation/compliance/pii-exports')
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        results = response.data.get('results', response.data)
        self.assertEqual(len(results), 2)
        # Ordered by -created_at
        self.assertEqual(results[0]['id'], self.log_a2.id)
        self.assertEqual(results[1]['id'], self.log_a1.id)

    def test_filter_by_report_key(self):
        self.client.force_authenticate(user=self.admin_a)
        response = self.client.get('/api/v1/foundation/compliance/pii-exports', {'report_key': REPORT_KEY_DAPODIK})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        results = response.data.get('results', response.data)
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]['id'], self.log_a1.id)

    def test_filter_by_school_id(self):
        self.client.force_authenticate(user=self.admin_a)
        response = self.client.get('/api/v1/foundation/compliance/pii-exports', {'school_id': 102})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        results = response.data.get('results', response.data)
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]['id'], self.log_a2.id)

    def test_cross_tenant_isolation(self):
        """Foundation B admin must only see Foundation B's logs, never Foundation A."""
        self.client.force_authenticate(user=self.admin_b)
        response = self.client.get('/api/v1/foundation/compliance/pii-exports')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        results = response.data.get('results', response.data)
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]['id'], self.log_b1.id)

    def test_permission_denied_without_audit_log_read(self):
        self.client.force_authenticate(user=self.teacher_a)
        response = self.client.get('/api/v1/foundation/compliance/pii-exports')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

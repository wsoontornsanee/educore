"""DSAR access export (CMP-011), reusing the CMP-016 PII-export ExportJob pipeline."""
import datetime
import io
from decimal import Decimal
from unittest import mock

import openpyxl
from django.test import TestCase

from apps.academic.models import AcademicYear, ClassEnrollment, ClassGroup
from apps.attendance.models import AttendanceDay, AttendanceStatus
from apps.compliance.exports import REPORT_KEY_DSAR_ACCESS
from apps.compliance.models import PiiExportAccessLog
from apps.compliance.services import PersonNotFoundError, collect_person_data_bundle
from apps.core.models import ExportJob
from apps.core.services import (
    create_export_job,
    get_export_allowed_formats,
    get_export_renderer,
    is_export_pii,
    run_export_job,
)
from apps.finance.models import Invoice, InvoiceStatus
from apps.identity.models import Foundation, Person, School, Staff, Student, User
from educore.middleware.tenancy import clear_current_foundation_id, set_current_foundation_id


class CollectPersonDataBundleTests(TestCase):
    def setUp(self):
        clear_current_foundation_id()
        self.foundation = Foundation.objects.create(
            legal_name="Yayasan Uji", brand_name="Uji", npwp="01.000.000.0-002.000",
        )
        set_current_foundation_id(self.foundation.id)
        self.school = School.all_tenants.create(
            foundation_id=self.foundation.id, name="SMA Uji", npsn="40100099", level=School.LEVEL_SMA,
        )
        self.person = Person.all_tenants.create(
            foundation_id=self.foundation.id, full_name="Siti Aminah",
            nik="3171010101019999", dob=datetime.date(2009, 1, 1),
        )
        self.student = Student.all_tenants.create(
            foundation_id=self.foundation.id, school=self.school, person=self.person,
            nis="2026099", nisn="0099999999", status=Student.STATUS_ACTIVE,
        )
        self.academic_year = AcademicYear.all_tenants.create(
            foundation_id=self.foundation.id, school=self.school,
            name="2026/2027", start_date=datetime.date(2026, 7, 1), end_date=datetime.date(2027, 6, 30),
        )
        self.rombel = ClassGroup.all_tenants.create(
            foundation_id=self.foundation.id, school=self.school,
            academic_year=self.academic_year, grade_level=10, name="X IPA 1",
        )
        ClassEnrollment.all_tenants.create(
            foundation_id=self.foundation.id, student=self.student,
            class_group=self.rombel, enrolled_at=datetime.date(2026, 7, 15), is_active=True,
        )
        AttendanceDay.all_tenants.create(
            foundation_id=self.foundation.id, school=self.school, student=self.student,
            date=datetime.date(2026, 8, 3), status=AttendanceStatus.HADIR,
        )
        Invoice.objects.create(
            foundation_id=self.foundation.id, school=self.school, student=self.student,
            number="INV/UJI/2026/000099", period="2026-08", due_date=datetime.date(2026, 8, 10),
            subtotal=Decimal('1000000.00'), total=Decimal('1000000.00'), currency='IDR',
            status=InvoiceStatus.ISSUED,
        )

    def tearDown(self):
        clear_current_foundation_id()

    def test_bundle_contains_all_four_domains(self):
        bundle = collect_person_data_bundle('STUDENT', self.student.id, self.foundation.id)
        self.assertEqual(bundle['identity']['full_name'], 'Siti Aminah')
        self.assertEqual(bundle['identity']['nik'], '3171010101019999')
        self.assertEqual(len(bundle['academic_enrollments']), 1)
        self.assertEqual(bundle['academic_enrollments'][0]['class_group'], 'X IPA 1')
        self.assertEqual(len(bundle['attendance_days']), 1)
        self.assertEqual(bundle['attendance_days'][0]['status'], 'HADIR')
        self.assertEqual(len(bundle['invoices']), 1)
        self.assertEqual(bundle['invoices'][0]['number'], 'INV/UJI/2026/000099')

    def test_unknown_student_raises(self):
        with self.assertRaises(PersonNotFoundError):
            collect_person_data_bundle('STUDENT', 999999, self.foundation.id)


class DsarAccessExportRegistrationTests(TestCase):
    """CMP-016 wiring: dsar_access must be PII-registered so watermark +
    PiiExportAccessLog fire automatically via run_export_job (unchanged from
    the DAPODIK/EMIS exporter's own registration)."""

    def test_renderer_registered(self):
        self.assertIsNotNone(get_export_renderer(REPORT_KEY_DSAR_ACCESS))

    def test_registered_as_pii(self):
        self.assertTrue(is_export_pii(REPORT_KEY_DSAR_ACCESS))

    def test_allowed_formats_xlsx_only(self):
        self.assertEqual(get_export_allowed_formats(REPORT_KEY_DSAR_ACCESS), {ExportJob.FORMAT_XLSX})


class DsarAccessExportRendererTests(TestCase):
    def setUp(self):
        clear_current_foundation_id()
        self.foundation = Foundation.objects.create(
            legal_name="Yayasan Uji Render", brand_name="Uji Render", npwp="01.000.000.0-007.000",
        )
        set_current_foundation_id(self.foundation.id)
        self.school = School.all_tenants.create(
            foundation_id=self.foundation.id, name="SMA Uji Render", npsn="40100093", level=School.LEVEL_SMA,
        )
        self.person = Person.all_tenants.create(
            foundation_id=self.foundation.id, full_name="Rendi Saputra", nik="3171010101014444",
        )
        self.student = Student.all_tenants.create(
            foundation_id=self.foundation.id, school=self.school, person=self.person,
            nis="2026093", status=Student.STATUS_ACTIVE,
        )

    def tearDown(self):
        clear_current_foundation_id()

    def test_renderer_produces_xlsx_bytes(self):
        renderer = get_export_renderer(REPORT_KEY_DSAR_ACCESS)
        job = ExportJob.objects.create(
            foundation_id=self.foundation.id,
            report_key=REPORT_KEY_DSAR_ACCESS,
            format=ExportJob.FORMAT_XLSX,
            filters={'subject_type': 'STUDENT', 'subject_id': self.student.id},
        )
        data, content_type, filename = renderer(job)
        self.assertTrue(len(data) > 0)
        self.assertIn('spreadsheetml', content_type)
        self.assertTrue(filename.endswith('.xlsx'))

    def test_renderer_unknown_subject_raises(self):
        from apps.compliance.services import PersonNotFoundError

        renderer = get_export_renderer(REPORT_KEY_DSAR_ACCESS)
        job = ExportJob.objects.create(
            foundation_id=self.foundation.id,
            report_key=REPORT_KEY_DSAR_ACCESS,
            format=ExportJob.FORMAT_XLSX,
            filters={'subject_type': 'STUDENT', 'subject_id': 999999},
        )
        with self.assertRaises(PersonNotFoundError):
            renderer(job)


class CollectStaffDataBundleTests(TestCase):
    """The STAFF branch of the DSAR bundle collector (CMP-011)."""

    def setUp(self):
        clear_current_foundation_id()
        self.foundation = Foundation.objects.create(
            legal_name="Yayasan Uji Staf", brand_name="Uji Staf", npwp="01.000.000.0-008.000",
        )
        set_current_foundation_id(self.foundation.id)
        self.school = School.all_tenants.create(
            foundation_id=self.foundation.id, name="SMA Uji Staf", npsn="40100092", level=School.LEVEL_SMA,
        )
        self.person = Person.all_tenants.create(
            foundation_id=self.foundation.id, full_name="Budi Pengajar",
            nik="3171010101013333", dob=datetime.date(1985, 4, 2),
        )
        self.user = User.all_tenants.create_user(
            phone_e164="+6281777777771", foundation_id=self.foundation.id, full_name="Budi Pengajar",
        )
        self.staff = Staff.all_tenants.create(
            foundation_id=self.foundation.id, person=self.person, user=self.user, school=self.school,
            nip="198504022010011001", nuptk="1612345678900002",
            join_date=datetime.date(2010, 1, 1), status=Staff.STATUS_ACTIVE,
        )

    def tearDown(self):
        clear_current_foundation_id()

    def test_staff_bundle_shape(self):
        bundle = collect_person_data_bundle('STAFF', self.staff.id, self.foundation.id)

        self.assertEqual(bundle['identity']['subject_type'], 'STAFF')
        self.assertEqual(bundle['identity']['full_name'], 'Budi Pengajar')
        self.assertEqual(bundle['identity']['nik'], '3171010101013333')
        self.assertEqual(bundle['identity']['nip'], '198504022010011001')
        self.assertEqual(bundle['identity']['nuptk'], '1612345678900002')
        self.assertEqual(bundle['identity']['status'], Staff.STATUS_ACTIVE)
        self.assertEqual(bundle['identity']['dob'], '1985-04-02')

        # Student-only domains are present but empty for staff subjects.
        for key in (
            'academic_enrollments', 'academic_report_cards',
            'attendance_days', 'period_attendances', 'invoices', 'payments',
        ):
            self.assertEqual(bundle[key], [], key)

    def test_unknown_staff_raises(self):
        with self.assertRaises(PersonNotFoundError):
            collect_person_data_bundle('STAFF', 999999, self.foundation.id)

    def test_staff_renderer_produces_xlsx(self):
        renderer = get_export_renderer(REPORT_KEY_DSAR_ACCESS)
        job = ExportJob.objects.create(
            foundation_id=self.foundation.id,
            report_key=REPORT_KEY_DSAR_ACCESS,
            format=ExportJob.FORMAT_XLSX,
            filters={'subject_type': 'STAFF', 'subject_id': self.staff.id},
        )
        data, content_type, filename = renderer(job)
        self.assertTrue(len(data) > 0)
        self.assertIn('spreadsheetml', content_type)
        self.assertIn(f'dsar_staff_{self.staff.id}_', filename)


class DsarAccessExportEndToEndTests(CollectPersonDataBundleTests):
    """Drive a dsar_access job through the whole CMP-016 pipeline: renderer ->
    is_export_pii -> watermark_export_data -> registered PII logger. Regression
    for the identity sheet having been titled 'Info', which the watermarker
    excludes from its record count (every DSAR log row read record_count=0)."""

    def setUp(self):
        super().setUp()
        self.admin_user = User.all_tenants.create_user(
            phone_e164="+6281777777772", foundation_id=self.foundation.id, full_name="Admin DSAR",
        )

    @mock.patch('apps.core.storage.upload_bytes')
    def test_pii_access_log_records_real_record_count(self, mock_upload):
        uploaded = []
        mock_upload.side_effect = lambda key, data, content_type: uploaded.append((key, data, content_type))

        job = create_export_job(
            report_key=REPORT_KEY_DSAR_ACCESS,
            export_format=ExportJob.FORMAT_XLSX,
            filters={'subject_type': 'STUDENT', 'subject_id': self.student.id},
            foundation_id=self.foundation.id,
            requested_by=str(self.admin_user.id),
            requested_by_name=self.admin_user.full_name,
        )
        run_export_job({'export_job_id': job.id})

        job.refresh_from_db()
        self.assertEqual(job.status, ExportJob.STATUS_COMPLETED)

        log = PiiExportAccessLog.objects.filter(export_job=job).first()
        self.assertIsNotNone(log)
        self.assertEqual(log.report_key, REPORT_KEY_DSAR_ACCESS)
        self.assertEqual(log.exported_by_id, str(self.admin_user.id))
        self.assertTrue(log.watermark_text)
        # The fixture subject has 1 enrollment + 1 attendance day + 1 invoice
        # (3 data rows). The identity sheet must count too — it did not while
        # it was titled 'Info', which the watermarker excludes.
        self.assertGreater(log.record_count, 3)

    @mock.patch('apps.core.storage.upload_bytes')
    def test_identity_sheet_is_separate_from_watermark_info_sheet(self, mock_upload):
        uploaded = []
        mock_upload.side_effect = lambda key, data, content_type: uploaded.append((key, data, content_type))

        job = create_export_job(
            report_key=REPORT_KEY_DSAR_ACCESS,
            export_format=ExportJob.FORMAT_XLSX,
            filters={'subject_type': 'STUDENT', 'subject_id': self.student.id},
            foundation_id=self.foundation.id,
            requested_by=str(self.admin_user.id),
            requested_by_name=self.admin_user.full_name,
        )
        run_export_job({'export_job_id': job.id})

        _key, data, _content_type = uploaded[0]
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertIn('Identity', wb.sheetnames)

        identity_text = " ".join(
            str(cell.value) for row in wb['Identity'].iter_rows() for cell in row if cell.value
        )
        self.assertIn('Siti Aminah', identity_text)
        # The watermarker's audit block lands in its own 'Info' sheet, not in
        # the subject's identity rows.
        self.assertNotIn('RAHASIA / PII', identity_text)
        self.assertIn('Info', wb.sheetnames)
        info_text = " ".join(
            str(cell.value) for row in wb['Info'].iter_rows() for cell in row if cell.value
        )
        self.assertIn('RAHASIA / PII', info_text)

"""Export pipeline wiring for statutory exports (CMP-018, RPT-002/003).

Registers DAPODIK/EMIS XLSX/CSV renderers with the core ExportJob
infrastructure. Imported from ComplianceConfig.ready() so registration
happens at startup.
"""
import csv
import io
import logging

from apps.compliance.models import PiiExportAccessLog, StatutorySystem
from apps.compliance.services import StatutoryExportError, get_exporter, validate_statutory_export
from apps.core.models import ExportJob
from apps.core.services import (
    register_export_formats,
    register_export_permission,
    register_export_pii,
    register_export_renderer,
    register_pii_export_logger,
)
from apps.identity.models import School

logger = logging.getLogger(__name__)

REPORT_KEY_DAPODIK = 'statutory_dapodik'
REPORT_KEY_EMIS = 'statutory_emis'

_SHEET_ORDER = ['students', 'staff', 'rombel']


def _get_school(foundation_id: int, school_id) -> School:
    school = School.objects.filter(
        foundation_id=foundation_id,
        id=school_id,
    ).first()
    if school is None:
        raise StatutoryExportError(f"School {school_id} not found in foundation {foundation_id}")
    return school


def _bundle_to_xlsx(bundle: dict) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)

    metadata = bundle.get('metadata', {})
    for sheet_name in _SHEET_ORDER:
        rows = bundle['sheets'].get(sheet_name, [])
        ws = wb.create_sheet(title=sheet_name[:31])
        if rows:
            columns = list(rows[0].keys())
            ws.append(columns)
            for row in rows:
                ws.append([row.get(col) for col in columns])
        else:
            ws.append(['(kosong / empty)'])

    info = wb.create_sheet(title='Info', index=0)
    info.append(['Sistem', metadata.get('system', '')])
    info.append(['Versi Skema', metadata.get('schema_version', '')])
    info.append(['Sekolah', metadata.get('school', '')])
    info.append(['NPSN', metadata.get('npsn', '')])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _bundle_to_csv(bundle: dict, sheet_name: str) -> bytes:
    rows = bundle['sheets'].get(sheet_name, [])
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    if rows:
        columns = list(rows[0].keys())
        writer.writerow(columns)
        for row in rows:
            writer.writerow([row.get(col) for col in columns])
    else:
        writer.writerow(['(kosong / empty)'])
    return buffer.getvalue().encode('utf-8-sig')


def _render_statutory_export(job: ExportJob):
    """Shared renderer for both statutory report keys.

    Filters: {'school_id': int, 'system': 'DAPODIK'|'EMIS'}. Returns
    XLSX bundle (one sheet per entity + Info) or per-sheet CSV.
    """
    filters = job.filters or {}
    school = _get_school(job.foundation_id, filters.get('school_id'))
    system = filters.get('system') or StatutorySystem.DAPODIK

    exporter = get_exporter(system, school)
    bundle = exporter.build_bundle()

    if job.format == ExportJob.FORMAT_XLSX:
        data = _bundle_to_xlsx(bundle)
        return data, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', f'statutory_{system.lower()}_{job.id}.xlsx'

    sheet = filters.get('sheet', 'students')
    if sheet not in _SHEET_ORDER:
        raise StatutoryExportError(f"Unknown sheet for CSV export: {sheet}")
    data = _bundle_to_csv(bundle, sheet)
    return data, 'text/csv', f'statutory_{system.lower()}_{sheet}_{job.id}.csv'


@register_export_renderer(REPORT_KEY_DAPODIK)
def render_dapodik_export(job: ExportJob):
    return _render_statutory_export(job)


@register_export_renderer(REPORT_KEY_EMIS)
def render_emis_export(job: ExportJob):
    return _render_statutory_export(job)


register_export_formats(REPORT_KEY_DAPODIK, {ExportJob.FORMAT_XLSX, ExportJob.FORMAT_CSV})
register_export_formats(REPORT_KEY_EMIS, {ExportJob.FORMAT_XLSX, ExportJob.FORMAT_CSV})
register_export_pii(REPORT_KEY_DAPODIK)
register_export_pii(REPORT_KEY_EMIS)


@register_pii_export_logger
def log_statutory_pii_export(job: ExportJob, filename: str, data: bytes, watermark_text: str, record_count: int):
    """Record immutable PII access log row for statutory exports (CMP-016, RPT-004)."""
    filters = job.filters or {}
    school_id = filters.get('school_id')
    try:
        school_id = int(school_id) if school_id is not None else None
    except (TypeError, ValueError):
        school_id = None

    PiiExportAccessLog.objects.create(
        foundation_id=job.foundation_id,
        export_job=job,
        report_key=job.report_key,
        format=job.format,
        exported_by_id=str(job.requested_by or ''),
        exported_by_name=str(job.requested_by_name or ''),
        school_id=school_id,
        filters=filters,
        record_count=record_count,
        watermark_text=watermark_text,
        file_name=filename,
        file_size=len(data),
    )

